"""Initial import from the spreadsheet XLSX (design §5), on a fake workbook.

The workbook copies the real sheet names and column headers; every value is synthetic and uses
the config.example persona.
"""

import sqlite3
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from lodestar.jobs.initial_import import main, run_import

ROOT = Path(__file__).parents[3]
EXAMPLE = ROOT / "config.example"

# Real header of "Log Entrate_Uscite": the per-category subcategory columns are form plumbing.
LOG_HEADER = [
    "Informazioni cronologiche", "Data", "Categoria",
    "💼 Stipendio", "🛒 Supermercato", "🍸 Cibo fuori / Bevute", "🚙 Auto", "🎾 Sport",
    "🏠 Casa", "👕 Abbigliamento", "🔁 Abbonamenti", "🎉 Divertimento", "📈 Investimenti",
    "🎁 Regali", "✈️ Viaggi", "💇 Cura / Benessere", "🚊 Trasporto pubblico", "⚕️ Salute",
    "🎭 Cultura", "📦 Altro", "💻 Tecnologia", "🏠 Affitto",
    "Sottocategoria aggiuntiva 1", "Sottocategoria aggiuntiva 2", "Sottocategoria aggiuntiva 3",
    "Sottocategoria", "Importo per pivot E/U", "Importo", "Note", "Da", "A",
    "Tracciamento in Entrate/Uscite", "Tag",
]  # fmt: skip
TRADES_HEADER = [
    "Data", "Ticker", "Quantità", "Prezzo d'acquisto", "Commissioni", "Valore di carico",
    "Totale", None,
    "Ticker", "Asset", "Prezzo - Metodo 1 (G. Finance)", "Prezzo - Metodo 2 (ISIN)",
    "Prezzo - Metodo 3 (URL)", "Tassazione", "Quantità", "Prezzo attuale", "Helper 1", "Helper 2",
]  # fmt: skip


def log_row(
    stamp: str, day: str, category: str, sub: str, amount: object, note: str,
    src: str, dst: str, tracking: str = "", tag: str = "",
) -> list[object]:  # fmt: skip
    row: list[object] = [None] * len(LOG_HEADER)
    values = {
        "Informazioni cronologiche": stamp, "Data": day, "Categoria": category,
        "Sottocategoria": sub, "Importo": amount, "Note": note, "Da": src, "A": dst,
        "Tracciamento in Entrate/Uscite": tracking, "Tag": tag,
    }  # fmt: skip
    for name, value in values.items():
        row[LOG_HEADER.index(name)] = value
    return row


def build_sheet(path: Path) -> Path:
    wb = Workbook()
    log = wb.active
    assert log is not None
    log.title = "Log Entrate_Uscite"
    log.append(LOG_HEADER)
    log.append(log_row("05/03/2026 09.10.00", "05/03/2026", "💰 Stipendio", "", "€ 2.100,00",
                       "Stipendio marzo", "", "Banca Aurora"))  # fmt: skip
    log.append(log_row("06/03/2026 21.15.30", "06/03/2026", "🍝 Ristoranti", "Pizzeria",
                       "€ 18,50", "Pizza", "Contanti", "", tag="cena-amici"))  # fmt: skip
    log.append([None] * len(LOG_HEADER))  # blank rows are skipped
    log.append(log_row("29/07/2026 18.00.00", "29/07/2026", "📊 Investimenti", "ETF",
                       "€ 1.400,00", "Acquisto XEON", "Broker Orsa", "", "No"))  # fmt: skip
    # Typed cells, as an XLSX downloaded from Google Sheets may have them.
    log.append(log_row(datetime(2026, 7, 1, 8, 0), datetime(2026, 7, 1), "❓ Varie", "",
                       650.0, "Giroconto", "Banca Aurora", "Broker Orsa"))  # fmt: skip

    trades = wb.create_sheet("Log Transazioni")
    trades.append(["Log acquisti e vendite", *[None] * 7, "Lista posizioni aperte e chiuse"])
    trades.append(TRADES_HEADER)
    trades.append(["29/07/2026", "XEON", "10,000000", "€ 140,00", "€ 0,00", "", "", None,
                   "XEON", "ETF Monetario", "", "LU0000000XE1", "", "26,0%"])  # fmt: skip
    trades.append(["01/08/2026", "BTP-01GN31", "5,000000", "€ 98,5000", "€ 2,95", "", "", None,
                   "BTP-01GN31", "Titolo di Stato", "", "IT0000000AA1", "", "12,5%"])  # fmt: skip
    trades.append(["02/08/2026", "BTC", "0,010000", "€ 60.000,00", "€ 1,00", "", "", None,
                   "BTC", "Crypto", "CURRENCY:BTC-EUR", "", "", "26,0%"])  # fmt: skip
    trades.append([None] * 8 + ["VWCE", "ETF Azionario", "", "IE00000VWCE1", "", "26,0%"])

    for name in ("E_U", "Patrimonio", "Dati e Calcoli"):
        wb.create_sheet(name).append(["computed view, never read"])
    wb.save(path)
    return path


@pytest.fixture
def sheet(tmp_path: Path) -> Path:
    return build_sheet(tmp_path / "Situazione patrimoniale.xlsx")


def rows(data_dir: Path, sql: str) -> list[tuple[object, ...]]:
    with sqlite3.connect(data_dir / "lodestar.db") as conn:
        return conn.execute(sql).fetchall()


def test_log_rows_become_transactions(sheet: Path, tmp_path: Path) -> None:
    data = tmp_path / "data"
    outcome = run_import(sheet, EXAMPLE, data, dry_run=False)
    assert outcome.errors == []
    imported = rows(
        data,
        "SELECT date, amount_cents, from_account, to_account, category, subcategory, note, tag,"
        " exclude_from_stats FROM transactions WHERE source = 'import' ORDER BY date",
    )
    assert imported == [
        ("2026-03-05", 210000, None, "Banca Aurora", "💰 Stipendio", None, "Stipendio marzo",
         None, 0),
        ("2026-03-06", 1850, "Contanti", None, "🍝 Ristoranti", "Pizzeria", "Pizza",
         "cena-amici", 0),
        ("2026-07-01", 65000, "Banca Aurora", "Broker Orsa", "❓ Varie", None, "Giroconto",
         None, 0),
        ("2026-07-29", 140000, "Broker Orsa", None, "📊 Investimenti", "ETF", "Acquisto XEON",
         None, 1),
    ]  # fmt: skip


def test_timestamp_is_kept_as_created_at_in_utc(sheet: Path, tmp_path: Path) -> None:
    data = tmp_path / "data"
    run_import(sheet, EXAMPLE, data, dry_run=False)
    [(created,)] = rows(data, "SELECT created_at FROM transactions WHERE date = '2026-03-06'")
    assert str(created).startswith("2026-03-06 20:15:30")  # 21:15:30 in Rome, CET


def test_plan_adjustment_is_added(sheet: Path, tmp_path: Path) -> None:
    data = tmp_path / "data"
    run_import(sheet, EXAMPLE, data, dry_run=False)
    adjustments = rows(
        data,
        "SELECT date, amount_cents, from_account, to_account, category, note,"
        " exclude_from_stats FROM transactions WHERE source = 'adjustment'",
    )
    assert adjustments == [
        ("2026-03-01", 123456, None, "Banca Aurora", "❓ Varie", "Riallineamento saldi", 1)
    ]


def test_trades_and_instruments(sheet: Path, tmp_path: Path) -> None:
    data = tmp_path / "data"
    run_import(sheet, EXAMPLE, data, dry_run=False)
    instruments = rows(
        data, "SELECT ticker, isin, asset_class, tax_rate, bucket FROM instruments ORDER BY ticker"
    )
    assert instruments == [
        ("BTP-01GN31", "IT0000000AA1", "govt_bond", "0.125", "P3"),
        ("VWCE", "IE00000VWCE1", "equity_etf", "0.26", "P4"),  # watchlist: no trade
        ("XEON", "LU0000000XE1", "money_market_etf", "0.26", "P2"),
    ]
    trades = rows(
        data,
        "SELECT t.date, i.ticker, t.side, t.quantity, t.price, t.fees_cents, t.account,"
        " tx.note FROM trades t JOIN instruments i ON i.id = t.instrument_id"
        " LEFT JOIN transactions tx ON tx.id = t.transaction_id ORDER BY t.date",
    )
    assert trades == [
        ("2026-07-29", "XEON", "buy", "10", "140", 0, "Broker Orsa", "Acquisto XEON"),
        # Sheet quantity is nominal ÷ 100: 5 → 500 nominal, price stays % of nominal.
        ("2026-08-01", "BTP-01GN31", "buy", "500", "98.5", 295, "Broker Orsa", None),
    ]


def test_dry_run_reports_counts_and_skips(sheet: Path, tmp_path: Path) -> None:
    data = tmp_path / "data"
    outcome = run_import(sheet, EXAMPLE, data, dry_run=True)
    assert outcome.summary.count == 5  # 4 log rows + the adjustment
    report = "\n".join(outcome.notes)
    assert "Strumenti: 3 (in watchlist: 1)" in report
    assert "Operazioni: 2 (collegate a un movimento: 1)" in report
    assert "BTC" in report  # skipped: no ISIN
    assert rows(data, "SELECT count(*) FROM transactions") == [(0,)]
    assert rows(data, "SELECT count(*) FROM trades") == [(0,)]


def test_invalid_row_aborts_everything(tmp_path: Path) -> None:
    path = build_sheet(tmp_path / "bad.xlsx")
    from openpyxl import load_workbook

    wb = load_workbook(path)
    wb["Log Entrate_Uscite"].append(
        log_row("", "07/03/2026", "Categoria inventata", "", "€ 5,00", "", "Contanti", "")
    )
    wb.save(path)
    data = tmp_path / "data"
    outcome = run_import(path, EXAMPLE, data, dry_run=False)
    assert any("Log Entrate_Uscite" in e and "riga 7" in e for e in outcome.errors)
    assert rows(data, "SELECT count(*) FROM transactions") == [(0,)]
    assert rows(data, "SELECT count(*) FROM instruments") == [(0,)]


def test_make_import_accepts_xlsx(
    sheet: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    code = main([str(sheet), "--dry-run", "--config", str(EXAMPLE), "--data", str(tmp_path / "d")])
    assert code == 0
    assert "Operazioni: 2" in capsys.readouterr().out


def test_dates_are_read(sheet: Path, tmp_path: Path) -> None:
    outcome = run_import(sheet, EXAMPLE, tmp_path / "data", dry_run=True)
    assert outcome.summary.first_date == date(2026, 3, 1)


@pytest.mark.parametrize("title", ["Log Entrate-Uscite", "log entrate uscite", "LogEntrateUscite"])
def test_sheet_names_match_ignoring_spacing_and_punctuation(tmp_path: Path, title: str) -> None:
    from openpyxl import load_workbook

    path = build_sheet(tmp_path / "s.xlsx")
    wb = load_workbook(path)
    wb["Log Entrate_Uscite"].title = title
    wb.save(path)
    outcome = run_import(path, EXAMPLE, tmp_path / "data", dry_run=True)
    assert outcome.errors == []


def test_missing_sheet_error_lists_the_sheets_found(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    path = build_sheet(tmp_path / "s.xlsx")
    wb = load_workbook(path)
    wb["Log Entrate_Uscite"].title = "Movimenti 2026"
    wb.save(path)
    [error] = run_import(path, EXAMPLE, tmp_path / "data", dry_run=True).errors
    assert "'Movimenti 2026'" in error and "'Log Transazioni'" in error
