import csv
import io
import sqlite3
from datetime import date, datetime
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from lodestar.adapters.loaders.csv_loader import HEADER
from lodestar.jobs.initial_import import run_import

ROOT = Path(__file__).parents[3]
EXAMPLE = ROOT / "config.example"
FAKE_LEDGER = EXAMPLE / "ledger.csv"
URL = "/api/transactions/export"


@pytest.fixture
def client(app: FastAPI) -> TestClient:
    run_import(FAKE_LEDGER, EXAMPLE, app.state.db_path.parent, dry_run=False)
    return TestClient(app)


def ledger(db: Path) -> list[tuple[object, ...]]:
    with sqlite3.connect(db) as conn:
        return conn.execute(
            "SELECT date, amount_cents, from_account, to_account, category, subcategory, note,"
            " tag, exclude_from_stats FROM transactions"
            " ORDER BY date, amount_cents, category, note"
        ).fetchall()


def test_csv_export_reimports_to_the_same_ledger(client: TestClient, tmp_path: Path) -> None:
    response = client.get(URL, params={"format": "csv"})
    assert response.status_code == 200
    exported = tmp_path / "export.csv"
    exported.write_bytes(response.content)

    outcome = run_import(exported, EXAMPLE, tmp_path / "fresh", dry_run=False)
    assert outcome.errors == []
    original = client.app.state.db_path  # type: ignore[attr-defined]
    assert ledger(tmp_path / "fresh" / "lodestar.db") == ledger(original)


def test_csv_layout_and_download_headers(client: TestClient) -> None:
    response = client.get(URL, params={"format": "csv"})
    assert response.headers["content-type"].startswith("text/csv")
    disposition = response.headers["content-disposition"]
    assert "attachment" in disposition
    assert f"lodestar-movimenti-{date.today().isoformat()}.csv" in disposition
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8"))))
    assert rows[0] == HEADER
    assert rows[1][0] <= rows[-1][0]  # oldest first
    assert len(rows) - 1 == len(FAKE_LEDGER.read_text("utf-8").splitlines()) - 1


def test_export_applies_the_filters(client: TestClient) -> None:
    response = client.get(URL, params={"format": "csv", "account": "Contanti"})
    rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8"))))
    assert rows
    assert all("Contanti" in (r["from_account"], r["to_account"]) for r in rows)


def test_xlsx_has_real_numbers_and_dates(client: TestClient) -> None:
    response = client.get(URL, params={"format": "xlsx", "tag": "vacanza-lisbona-2026"})
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert ".xlsx" in response.headers["content-disposition"]

    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet is not None
    assert sheet.title == "Movimenti"
    values = list(sheet.iter_rows(values_only=True))
    assert list(values[0]) == HEADER
    first = values[1]
    assert isinstance(first[0], datetime)
    assert isinstance(first[1], float | int)
    assert first[8] in (True, False)
    assert len(values) - 1 == 8  # the Lisbon trip rows in the fake ledger
    assert sheet.freeze_panes == "A2"


def test_unknown_format_rejected(client: TestClient) -> None:
    assert client.get(URL, params={"format": "pdf"}).status_code == 422
