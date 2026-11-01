"""XLSX loader for the user's spreadsheet (design §5, XLSX import mapping).

Reads two sheets and ignores the rest (computed views, examples):
- "Log Entrate_Uscite" → ledger lines. The per-category subcategory columns are form plumbing;
  the combined "Sottocategoria" column holds the value.
- "Log Transazioni" → trades (left table) and instruments (right table, "Lista posizioni").

Cells may be typed (numbers, dates) or Italian text ("€ 1.234,56", "29/10/2026"); both parse.
"""

import contextlib
import re
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from lodestar.adapters.loaders.base import LoadResult, RowError
from lodestar.config.models import SettingsFile
from lodestar.domain.errors import FieldError
from lodestar.domain.ledger import Catalog, LineInput, validate_line
from lodestar.domain.money import quantize
from lodestar.domain.portfolio import BOND_UNIT, InstrumentInput, TradeInput

LOG_SHEET = "Log Entrate_Uscite"
TRADES_SHEET = "Log Transazioni"
LOG_COLUMNS = (
    "Informazioni cronologiche",
    "Data",
    "Categoria",
    "Sottocategoria",
    "Importo",
    "Note",
    "Da",
    "A",
    "Tracciamento in Entrate/Uscite",
    "Tag",
)
TRADE_COLUMNS = ("Data", "Ticker", "Quantità", "Prezzo d'acquisto", "Commissioni")
POSITION_COLUMNS = ("Ticker", "Asset", "Prezzo - Metodo 2 (ISIN)", "Tassazione")
ISIN = re.compile(r"^[A-Z]{2}[A-Z0-9]{9}[0-9]$")
DEFAULT_TAX = {"govt_bond": Decimal("0.125")}
ETF_TAX = Decimal("0.26")


class _Bad(Exception):
    def __init__(self, field: str, message: str) -> None:
        super().__init__(message)
        self.field = field


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _date(value: Any, field: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(_text(value), "%d/%m/%Y").date()
    except ValueError:
        raise _Bad(field, f"{_text(value)!r} is not a date like 29/10/2026") from None


def _timestamp(value: Any, tz: ZoneInfo) -> datetime | None:
    if isinstance(value, datetime):
        local = value
    elif not _text(value):
        return None
    else:
        try:
            local = datetime.strptime(_text(value), "%d/%m/%Y %H.%M.%S")
        except ValueError:
            return None  # a missing form timestamp is not worth rejecting the row
    return local.replace(tzinfo=tz).astimezone(UTC)


def _number(value: Any, field: str) -> Decimal:
    if isinstance(value, bool):
        raise _Bad(field, f"{value!r} is not a number")
    if isinstance(value, int | float):
        return Decimal(str(value))
    text = _text(value).replace("€", "").replace("%", "").replace(" ", "").replace(" ", "")
    try:
        return Decimal(text.replace(".", "").replace(",", "."))
    except InvalidOperation:
        raise _Bad(field, f"{_text(value)!r} is not a number like 1.234,56") from None


def _plain(value: Decimal) -> Decimal:
    """Drop trailing zeros without switching to exponent notation."""
    return Decimal(format(value.normalize(), "f"))


def _key(title: str) -> str:
    """Compare sheet names on letters and digits: "Log Entrate/Uscite" = "log entrate_uscite"."""
    return "".join(ch for ch in title.casefold() if ch.isalnum())


def _sheet(workbook: Any, name: str) -> Worksheet | None:
    for sheet in workbook.worksheets:
        if _key(sheet.title) == _key(name):
            return sheet  # type: ignore[no-any-return]
    return None


def _columns(header: list[Any], names: tuple[str, ...], start: int = 0) -> dict[str, int] | str:
    cells = [_text(h) for h in header]
    found: dict[str, int] = {}
    for name in names:
        try:
            found[name] = cells.index(name, start)
        except ValueError:
            return name
    return found


class XlsxLoader:
    def __init__(self, settings: SettingsFile) -> None:
        self._settings = settings
        self._tz = ZoneInfo(settings.timezone)

    def load(self, path: Path, catalog: Catalog) -> LoadResult:
        result = LoadResult()
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            log = _sheet(workbook, LOG_SHEET)
            if log is None:
                found = ", ".join(repr(sheet.title) for sheet in workbook.worksheets)
                message = f"sheet {LOG_SHEET!r} not found; sheets in the file: {found}"
                result.errors.append(RowError(1, None, message))
                return result
            self._ledger(log, catalog, result)
            trades = _sheet(workbook, TRADES_SHEET)
            if trades is not None:
                self._portfolio(trades, result)
        finally:
            workbook.close()
        self._adjustment(catalog, result)
        return result

    def _ledger(self, sheet: Worksheet, catalog: Catalog, result: LoadResult) -> None:
        rows = list(sheet.iter_rows(values_only=True))
        cols = _columns(list(rows[0]) if rows else [], LOG_COLUMNS)
        if isinstance(cols, str):
            result.errors.append(RowError(1, f"{LOG_SHEET} · {cols}", "column missing"))
            return
        # "📦Amazon" in the sheet is "📦 Amazon" in config: match without spaces.
        categories = {"".join(c.split()): c for c in catalog.categories}

        for number, row in enumerate(rows[1:], start=2):

            def cell(name: str, row: tuple[Any, ...] = row) -> Any:
                i = cols[name]
                return row[i] if i < len(row) else None

            if not _text(cell("Data")) and not _text(cell("Importo")):
                continue
            try:
                line = LineInput(
                    date=_date(cell("Data"), "Data"),
                    amount=quantize(_number(cell("Importo"), "Importo")),
                    from_account=_text(cell("Da")) or None,
                    to_account=_text(cell("A")) or None,
                    category=categories.get("".join(_text(cell("Categoria")).split()), "")
                    or _text(cell("Categoria")),
                    subcategory=_text(cell("Sottocategoria")) or None,
                    note=_text(cell("Note")) or None,
                    tag=_text(cell("Tag")) or None,
                    exclude_from_stats=_text(cell("Tracciamento in Entrate/Uscite")).lower()
                    == "no",
                )
                if line.amount == 0:  # placeholder rows, e.g. "opening" an account
                    result.skipped.append(f"{LOG_SHEET} riga {number}: importo 0, non importata")
                    continue
                validate_line(line, catalog)
            except (_Bad, FieldError) as exc:
                result.errors.append(RowError(number, f"{LOG_SHEET} · {exc.field}", str(exc)))
            else:
                result.rows.append(line)
                result.created_at.append(_timestamp(cell("Informazioni cronologiche"), self._tz))

    def _portfolio(self, sheet: Worksheet, result: LoadResult) -> None:
        rows = list(sheet.iter_rows(values_only=True))
        header = list(rows[1]) if len(rows) > 1 else []  # row 1 holds the two table titles
        trade_cols = _columns(header, TRADE_COLUMNS)
        if isinstance(trade_cols, str):
            result.errors.append(RowError(2, f"{TRADES_SHEET} · {trade_cols}", "column missing"))
            return
        position_cols = _columns(header, POSITION_COLUMNS, start=trade_cols["Ticker"] + 1)
        if isinstance(position_cols, str):
            result.errors.append(RowError(2, f"{TRADES_SHEET} · {position_cols}", "column missing"))
            return

        def at(row: tuple[Any, ...], i: int) -> Any:
            return row[i] if i < len(row) else None

        imp = self._settings.initial_import
        known: dict[str, InstrumentInput] = {}
        skipped: set[str] = set()
        for row in rows[2:]:
            ticker = _text(at(row, position_cols["Ticker"]))
            if not ticker:
                continue
            label = _text(at(row, position_cols["Asset"]))
            if ticker in imp.skip_tickers:
                result.skipped.append(f"{ticker}: escluso da configurazione")
                skipped.add(ticker)
                continue
            mapping = imp.instruments.get(ticker) or imp.asset_classes.get(label)
            if mapping is None:
                result.skipped.append(f"{ticker}: asset {label!r} non mappato")
                skipped.add(ticker)
                continue
            isin = (
                mapping.isin or _text(at(row, position_cols["Prezzo - Metodo 2 (ISIN)"]))
            ).upper()
            if not ISIN.match(isin):
                result.skipped.append(f"{ticker}: senza ISIN")
                skipped.add(ticker)
                continue
            tax_cell = at(row, position_cols["Tassazione"])
            if _text(tax_cell):
                tax = _number(tax_cell, "Tassazione")
                tax = tax / 100 if tax > 1 or "%" in _text(tax_cell) else tax
            else:
                tax = DEFAULT_TAX.get(mapping.asset_class, ETF_TAX)
            known[ticker] = InstrumentInput(
                ticker=mapping.ticker or ticker,
                isin=isin,
                name=mapping.ticker or ticker,
                asset_class=mapping.asset_class,
                tax_rate=_plain(tax),
                bucket=mapping.bucket,
            )
        traded = {_text(at(row, trade_cols["Ticker"])) for row in rows[2:]}
        for ticker, candidate in known.items():
            if imp.watchlist is None or ticker in traded or ticker in imp.watchlist:
                result.instruments.append(candidate)
            else:
                result.skipped.append(f"{ticker}: fuori watchlist")

        account = self._settings.portfolio.trade_account
        for number, row in enumerate(rows[2:], start=3):
            ticker = _text(at(row, trade_cols["Ticker"]))
            if not _text(at(row, trade_cols["Data"])) and not ticker:
                continue
            if ticker in skipped:
                when = _text(at(row, trade_cols["Data"]))
                with contextlib.suppress(_Bad):
                    when = _date(at(row, trade_cols["Data"]), "Data").strftime("%d/%m/%Y")
                result.skipped.append(f"{ticker}: operazione del {when} non importata")
                continue
            instrument = known.get(ticker)
            try:
                if instrument is None:
                    raise _Bad("Ticker", f"{ticker!r} is not in the positions list")
                quantity = _number(at(row, trade_cols["Quantità"]), "Quantità")
                if instrument.asset_class == "govt_bond":
                    quantity *= BOND_UNIT
                fees = at(row, trade_cols["Commissioni"])
                result.trades.append(
                    TradeInput(
                        date=_date(at(row, trade_cols["Data"]), "Data"),
                        ticker=instrument.ticker,
                        side="sell" if quantity < 0 else "buy",
                        quantity=_plain(abs(quantity)),
                        price=_plain(_number(at(row, trade_cols["Prezzo d'acquisto"]), "Prezzo")),
                        fees=quantize(_number(fees, "Commissioni")) if _text(fees) else Decimal(0),
                        account=account,
                    )
                )
            except _Bad as exc:
                result.errors.append(RowError(number, f"{TRADES_SHEET} · {exc.field}", str(exc)))

    def _adjustment(self, catalog: Catalog, result: LoadResult) -> None:
        adj = self._settings.initial_import.adjustment
        if adj is None:
            return
        line = LineInput(
            date=adj.date,
            amount=adj.amount,
            from_account=adj.account if adj.direction == "out" else None,
            to_account=adj.account if adj.direction == "in" else None,
            category=adj.category,
            subcategory=None,
            note=adj.note,
            tag=None,
            exclude_from_stats=True,
        )
        try:
            validate_line(line, catalog)
        except FieldError as exc:
            result.errors.append(RowError(0, f"initial_import.adjustment · {exc.field}", str(exc)))
        else:
            result.adjustments.append(line)
