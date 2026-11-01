"""Ledger export. The CSV uses the initial-import layout (design §5), so it re-imports as is."""

import csv
import io
from collections.abc import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font

from lodestar.adapters.db.tables import Transaction
from lodestar.adapters.loaders.csv_loader import HEADER
from lodestar.domain.money import from_cents

XLSX_MONEY = '#,##0.00 "€"'
XLSX_DATE = "dd/mm/yyyy"
WIDTHS = {"A": 12, "B": 12, "C": 18, "D": 18, "E": 26, "F": 18, "G": 40, "H": 22, "I": 10}


def _cells(row: Transaction) -> list[str]:
    return [
        row.date.isoformat(),
        f"{from_cents(row.amount_cents):.2f}",
        row.from_account or "",
        row.to_account or "",
        row.category,
        row.subcategory or "",
        row.note or "",
        row.tag or "",
        "true" if row.exclude_from_stats else "false",
    ]


def to_csv(rows: Iterable[Transaction]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(HEADER)
    writer.writerows(_cells(r) for r in rows)
    return buffer.getvalue().encode("utf-8")


def to_xlsx(rows: Iterable[Transaction]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Movimenti"
    sheet.append(HEADER)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for r in rows:
        sheet.append(
            [
                r.date,
                float(from_cents(r.amount_cents)),
                r.from_account,
                r.to_account,
                r.category,
                r.subcategory,
                r.note,
                r.tag,
                r.exclude_from_stats,
            ]
        )
        current = sheet.max_row
        sheet.cell(current, 1).number_format = XLSX_DATE
        sheet.cell(current, 2).number_format = XLSX_MONEY
    sheet.freeze_panes = "A2"
    for column, width in WIDTHS.items():
        sheet.column_dimensions[column].width = width
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
